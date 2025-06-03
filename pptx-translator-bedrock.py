import argparse
import boto3
import concurrent.futures
from pptx import Presentation

def translate_text(client, text, source_lang, target_lang, model_id):
    """Translate a single text segment using AWS Bedrock"""
    if not text.strip():
        return text
        
    try:
        response = client.converse(
            modelId=model_id,
            messages=[{
                "role": "user", 
                "content": [{"text": f"Translate from {source_lang} to {target_lang}: {text}"}]
            }]
        )
        
        content = response["output"]["message"]["content"]
        return next((item["text"] for item in content if "text" in item), str(content)) if isinstance(content, list) else str(content)
    except Exception as e:
        print(f"Translation error: {str(e)[:100]}...")
        return text

def process_slide(slide, client, source_lang, target_lang, model_id):
    """Process a single slide, translating all text elements"""
    # Process text frames
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text.strip():
                        run.text = translate_text(client, run.text, source_lang, target_lang, model_id)
        
        # Process tables
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for para in cell.text_frame.paragraphs:
                        for run in para.runs:
                            if run.text.strip():
                                run.text = translate_text(client, run.text, source_lang, target_lang, model_id)
    
    # Process slide notes
    if slide.has_notes_slide:
        for para in slide.notes_slide.notes_text_frame.paragraphs:
            for run in para.runs:
                if run.text.strip():
                    run.text = translate_text(client, run.text, source_lang, target_lang, model_id)
    
    return 1  # Return 1 for completed slide

def translate_pptx(input_file, source_lang, target_lang, model_id="us.amazon.nova-lite-v1:0", region=None, max_workers=4):
    """Translate a PowerPoint file using AWS Bedrock"""
    # Setup
    client = boto3.client("bedrock-runtime", region_name=region) if region else boto3.client("bedrock-runtime")
    prs = Presentation(input_file)
    total_slides = len(prs.slides)
    
    print(f"Translating {total_slides} slides from {source_lang} to {target_lang} using {max_workers} workers")
    
    # Translate slides in parallel
    with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
        # Map the function to all slides
        futures = {executor.submit(process_slide, slide, client, source_lang, target_lang, model_id): i 
                  for i, slide in enumerate(prs.slides)}
        
        # Process results as they complete
        completed = 0
        for future in concurrent.futures.as_completed(futures):
            completed += future.result()
            print(f"Progress: {completed}/{total_slides} slides processed")
    
    # Save the result
    output_file = f"{input_file.rsplit('.', 1)[0]}-{target_lang}.pptx"
    prs.save(output_file)
    
    print(f"\nTranslation complete!")
    print(f"Output file: {output_file}")
    
    return output_file

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="PowerPoint Translator using AWS Bedrock")
    parser.add_argument("source_lang", help="Source language code (e.g., en)")
    parser.add_argument("target_lang", help="Target language code (e.g., es)")
    parser.add_argument("file", help="PowerPoint file to translate")
    parser.add_argument("--model", default="us.amazon.nova-lite-v1:0", help="Bedrock model ID")
    parser.add_argument("--region", help="AWS region name")
    parser.add_argument("--workers", type=int, default=4, help="Number of parallel workers")
    args = parser.parse_args()
    
    translate_pptx(args.file, args.source_lang, args.target_lang, args.model, args.region, args.workers)
